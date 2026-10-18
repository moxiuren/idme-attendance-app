# -*- coding: utf-8 -*-
"""
1B 班干预报告排版与字号终极精修
显微镜级别对齐与字体字号显式规范
严禁 Emoji
"""
import sys, os, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')

TEMPLATE_PATH = r"C:\Users\Admin\Downloads\ORIGINAL_EMPTY_TEMPLATE.docx"
OUT_PATH = r"C:\Users\Admin\Downloads\1B Laporan Intervensi BMI 5-9T.docx"

doc = docx.Document(TEMPLATE_PATH)

# 元数据锁定
core_props = doc.core_properties
core_props.author = "GAN KAI PIN"
core_props.last_modified_by = "GAN KAI PIN"
core_props.title = "LAPORAN INTERVENSI PROGRAM BMI 5-9T 1B TAHUN 2026"
core_props.comments = ""

BLACK = RGBColor(0, 0, 0)

def set_run_style(run, font_name="Arial", size_pt=9.5, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK

def format_cell(cell, text, font_name="Arial", size_pt=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = cell.add_paragraph() if idx > 0 else cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        set_run_style(run, font_name=font_name, size_pt=size_pt, bold=bold)

# 1. 顶部标题排版
for p in doc.paragraphs:
    txt = p.text.strip()
    if "SJK (C) PEI HWA" in txt and "PERLING" in txt:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=12, bold=True)
    elif "LAPORAN INTERVENSI" in txt:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=True)

# 2. Table 1: 教师与班级
t1 = doc.tables[0]
format_cell(t1.rows[0].cells[0], "NAMA GURU:", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
format_cell(t1.rows[0].cells[1], "KELAS:", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
format_cell(t1.rows[0].cells[2], "INTERVENSI:", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

format_cell(t1.rows[1].cells[0], "GAN KAI PIN", size_pt=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
format_cell(t1.rows[1].cells[1], "1B", size_pt=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
format_cell(t1.rows[1].cells[2], "PENGGAL 1", size_pt=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

# 3. Table 2: 核心表格内容
t2 = doc.tables[1]

# Row 0: Tempoh Pelaksanaan
format_cell(t2.rows[0].cells[0], "Tempoh Pelaksanaan:", size_pt=9.5, bold=True)
format_cell(t2.rows[0].cells[1], "April hingga Julai 2026", size_pt=9.5, bold=False)

# Row 1: Hari/ Masa
format_cell(t2.rows[1].cells[0], "Hari/ Masa:", size_pt=9.5, bold=True)
format_cell(t2.rows[1].cells[1], "Hari Selasa, 4:50 petang - 5:20 petang\nHari Jumaat, 3:50 petang - 4:20 petang", size_pt=9.5, bold=False)

# Row 2: Objektif
format_cell(t2.rows[2].cells[0], "Objektif:", size_pt=9.5, bold=True)
format_cell(t2.rows[2].cells[1], "Memberi pendedahan kepada murid tentang amalan gaya hidup sihat, pemakanan seimbang dan latihan kecergasan fizikal bagi mencapai berat badan yang unggul.", size_pt=9.5, bold=False)

# Row 3: Sasaran (Kategori)
format_cell(t2.rows[3].cells[0], "Sasaran:", size_pt=9.5, bold=True)
format_cell(t2.rows[3].cells[1], "Murid Obes dan Kurus", size_pt=9.5, bold=False)

# Row 4: Murid Terlibat (纯英文名单，无多余空格)
format_cell(t2.rows[4].cells[0], "Murid Terlibat:", size_pt=9.5, bold=True)
format_cell(t2.rows[4].cells[1], "1. NUR AISYAH BALQIS BINTI AHMAD (Obes)\n2. LEE CHONG HAN (Obes)\n3. FARAH ADIBAH (Kurus)", size_pt=9.5, bold=False)

# Row 5: Intervensi / Cadangan Aktiviti
format_cell(t2.rows[5].cells[0], "Intervensi/ Cadangan Aktiviti:", size_pt=9.5, bold=True)
format_cell(t2.rows[5].cells[1], "1. Latihan Sistem Tenaga (Kecergasan Kardiovaskular)\n2. Senaman ringan dan Regangan Dinamik\n3. Berjalan pantas & Joging setempat 5 minit\n4. Bimbingan Pemakanan Sihat (Piramid Makanan KKM)", size_pt=9.5, bold=False)

# Row 6: Proses Pelaksanaan
format_cell(t2.rows[6].cells[0], "Proses Pelaksanaan:", size_pt=9.5, bold=True)
format_cell(t2.rows[6].cells[1], "1. Murid melakukan Latihan Sistem Tenaga yang ditetapkan secara berselang-seli semasa waktu PJ.\n2. Murid obes melakukan berjalan pantas dan joging setempat selama 5 minit sewaktu aktiviti pemanasan badan.\n3. Murid melakukan senaman lompat bintang sebanyak 30 kali untuk meningkatkan kadar metabolisme.\n4. Penerapan amalan pemakanan sihat serta galakan membawa bekal makanan berkhasiat ke sekolah.", size_pt=9.5, bold=False)

# Row 7: Kekuatan Intervensi
format_cell(t2.rows[7].cells[0], "Kekuatan Intervensi:", size_pt=9.5, bold=True)
format_cell(t2.rows[7].cells[1], "Intervensi yang dirancang dapat meningkatkan daya tahan kardiovaskular, tahap kecergasan fizikal dan membantu mengawal berat badan murid secara beransur-ansur.", size_pt=9.5, bold=False)

# Row 8: Kelemahan Intervensi
format_cell(t2.rows[8].cells[0], "Kelemahan Intervensi:", size_pt=9.5, bold=True)
format_cell(t2.rows[8].cells[1], "Guru sukar mengawal corak pemakanan dan aktiviti harian murid selain waktu PJ di sekolah.", size_pt=9.5, bold=False)

# Row 9: Dokumentasi Bergambar (纯净留白框)
format_cell(t2.rows[9].cells[0], "Dokumentasi Bergambar:", size_pt=9.5, bold=True)
c_img = t2.rows[9].cells[1]
c_img.text = ""
p_img = c_img.paragraphs[0]
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_before = Pt(30)
p_img.paragraph_format.space_after = Pt(30)
r_img = p_img.add_run("[ Sila tampal gambar aktiviti murid 1B sewaktu waktu PJ di sini ]")
set_run_style(r_img, font_name="Arial", size_pt=9, bold=False, italic=True)

# Row 10: Cadangan Penambahbaikan
format_cell(t2.rows[10].cells[0], "Cadangan Penambahbaikan:", size_pt=9.5, bold=True)
format_cell(t2.rows[10].cells[1], "1. Menukarkan senaman rutin kepada permainan kecil atau aktiviti senamrobik yang berunsur didik hibur.\n2. Memberi pujian, galakan dan ganjaran kecil sebagai motivasi kepada murid yang menunjukkan peningkatan kecergasan.\n3. Memberi kesedaran dan bimbingan ringkas kepada ibu bapa murid berkenaan tentang amalan pemakanan sihat di rumah.\n4. Merancang aktiviti intervensi lanjutan yang bersesuaian dengan tahap keupayaan murid pada Penggal 2.", size_pt=9.5, bold=False)

# 4. 签名处段落精准重构 (字号统一为 11pt, 字体全为 Arial, 消除多余空格与错位)
# 找到末尾签名段落并规范格式
for p in doc.paragraphs:
    txt = p.text.strip()
    if not txt:
        continue
    if "Disediakan oleh" in txt:
        p.text = "Disediakan oleh,\t\t\t\t\tDisemak oleh,"
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=False)
    elif "…………" in txt:
        p.text = "....................................................\t\t\t...................................................."
        p.paragraph_format.space_before = Pt(25)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=False)
    elif "LAU JEN SHAN" in txt:
        # 去掉 (GAN KAI PIN ) 后面的多余空格，变为规范的 (GAN KAI PIN)
        p.text = "(GAN KAI PIN)\t\t\t\t\t(LAU JEN SHAN)"
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=True)
    elif "GURU PJK" in txt:
        p.text = "GURU PJK\t\t\t\t\t\tKETUA PANITIA PJK"
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=False)
    elif "SJK(C) PEI HWA" in txt and p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        p.text = "SJK(C) PEI HWA\t\t\t\t\tSJK(C) PEI HWA"
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            set_run_style(r, font_name="Arial", size_pt=11, bold=False)

# 5. 全文档所有 run 强行保底锁定为纯黑 RGBColor(0,0,0)
for p in doc.paragraphs:
    for r in p.runs:
        r.font.color.rgb = BLACK

for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = BLACK

doc.save(OUT_PATH)
print(f"[精修完毕] 1B 班干预报告终极精修已完成: {OUT_PATH}")
